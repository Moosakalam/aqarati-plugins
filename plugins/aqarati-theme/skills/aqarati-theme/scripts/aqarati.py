#!/usr/bin/env python3
"""Aqarati theme: build PPTX/PDF from a JSON spec, or restyle an existing PPTX.

  python3 aqarati.py build spec.json out.pptx     # also writes out.pdf if --pdf
  python3 aqarati.py build spec.json out.pdf
  python3 aqarati.py restyle in.pptx out.pptx     # recolor/refont existing deck
"""
import json, sys, os, re, copy

HERE = os.path.dirname(os.path.abspath(__file__))
T = json.load(open(os.path.join(HERE, "..", "theme.json")))
C, F, SZ = T["colors"], T["fonts"], T["sizes_pt"]
W, H = T["slide"]["w_in"], T["slide"]["h_in"]
FOOTER = T["footer_text"]


def col(name):
    return C.get(name, name)  # theme key or raw hex


def badge_color(label):
    return col(T["status"].get(label.upper(), "green"))


# ───────────────────────── PPTX ─────────────────────────
def build_pptx(spec, out):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    rgb = lambda h: RGBColor.from_string(col(h))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    footer_text = spec.get("footer", FOOTER)

    def bg(s, h):
        f = s.background.fill; f.solid(); f.fore_color.rgb = rgb(h)

    def shape(s, kind, x, y, w, h, fill, radius=None, shadow=False):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background()
        if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        if not shadow:
            sh.shadow.inherit = False
        return sh

    def text(s, x, y, w, h, t, size, color, bold=False, align="l", anchor="t", font=None):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
        for i, line in enumerate(str(t).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            r = p.add_run(); r.text = line
            r.font.size = Pt(size); r.font.bold = bold
            r.font.name = font or F["main"]; r.font.color.rgb = rgb(color)
        return tb

    def pill(s, x, y, label, size=SZ["badge"]):
        w = max(0.9, 0.095 * len(label) + 0.45)
        shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.33, badge_color(label), radius=0.5)
        text(s, x, y, w, 0.33, label.upper(), size, "white", bold=True, align="c", anchor="m")
        return w

    def chrome(s, n, eyebrow, title, lede=None, badge=None):
        bg(s, "bg_light")
        text(s, 0.6, 0.45, 8, 0.35, eyebrow.upper(), SZ["eyebrow"], "green", bold=True)
        text(s, 0.6, 0.8, 10.5, 0.85, title, SZ["title"], "ink", bold=True)
        if lede:
            text(s, 0.6, 1.7, 11.5, 0.45, lede, SZ["lede"], "muted")
        if badge:
            pill(s, W - 0.6 - max(0.9, 0.095 * len(badge) + 0.45), 1.0, badge)
        text(s, 0.6, 7.07, 6, 0.3, footer_text, SZ["footer"], "muted")
        text(s, W - 1.6, 7.07, 1.0, 0.3, str(n), SZ["footer"], "muted", align="r")

    def card(s, x, y, w, h, title, body, tint=False, tag=None, icon=None):
        c = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, "green_tint" if tint else "card", radius=0.06, shadow=True)
        c.line.color.rgb = rgb("line"); c.line.width = Pt(0.75)
        d = 0.62
        shape(s, MSO_SHAPE.OVAL, x + 0.22, y + 0.22, d, d, "green_soft")
        if icon:
            text(s, x + 0.22, y + 0.22, d, d, icon, 18, "green", bold=True, align="c", anchor="m")
        if tag:
            pw = max(0.6, 0.085 * len(tag) + 0.35)
            shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + w - pw - 0.18, y + 0.24, pw, 0.27, badge_color(tag), radius=0.5)
            text(s, x + w - pw - 0.18, y + 0.24, pw, 0.27, tag.upper(), 9.5, "white", bold=True, align="c", anchor="m")
        text(s, x + 0.22, y + 0.98, w - 0.44, 0.4, title, SZ["card_title"] - 4, "ink", bold=True)
        text(s, x + 0.22, y + 1.42, w - 0.44, h - 1.5, body, SZ["card_body"] - 1.5, "body")

    for i, sl in enumerate(spec["slides"], 1):
        s = prs.slides.add_slide(blank)
        k = sl.get("type", "cards")
        if k == "cover":
            bg(s, "bg_dark")
            shape(s, MSO_SHAPE.OVAL, 9.4, -2.2, 5.6, 5.6, "circle_a")
            shape(s, MSO_SHAPE.OVAL, 10.2, 3.9, 4.6, 4.6, "circle_b")
            if sl.get("arabic"):
                text(s, 0.6, 1.35, 8, 0.6, sl["arabic"], SZ["cover_arabic"], "gold", bold=True)
            text(s, 0.6, 1.95, 9, 1.0, sl["title"], SZ["cover_title"], "white", bold=True)
            text(s, 0.6, 3.0, 9, 0.8, sl.get("subtitle", ""), SZ["cover_sub"] - 6, "lilac", bold=True)
            text(s, 0.6, 3.85, 8.5, 0.5, sl.get("tagline", ""), SZ["cover_tag"] - 2, "lilac_soft")
            text(s, 0.6, 5.6, 9, 0.8, sl.get("footnote", ""), 13, "lilac_dim")
        elif k == "section":
            bg(s, "bg_dark")
            shape(s, MSO_SHAPE.OVAL, 9.4, -2.2, 5.6, 5.6, "circle_a")
            text(s, 0.6, 2.6, 9, 0.4, sl.get("eyebrow", "").upper(), SZ["eyebrow"], "gold", bold=True)
            text(s, 0.6, 3.0, 10, 1.2, sl["title"], 44, "white", bold=True)
            if sl.get("lede"): text(s, 0.6, 4.3, 9, 0.6, sl["lede"], 18, "lilac")
        elif k == "bullets":
            chrome(s, i, sl.get("eyebrow", ""), sl["title"], sl.get("lede"), sl.get("badge"))
            y = 2.4
            for b in sl["bullets"]:
                shape(s, MSO_SHAPE.OVAL, 0.65, y + 0.1, 0.16, 0.16, "green")
                text(s, 1.0, y, 11, 0.5, b, 18, "body"); y += 0.62
        else:  # cards
            chrome(s, i, sl.get("eyebrow", ""), sl["title"], sl.get("lede"), sl.get("badge"))
            cards = sl["cards"]; n = len(cards)
            per = sl.get("per_row", min(n, 4)); rows = -(-n // per)
            gap = 0.3; top = 2.3 if sl.get("lede") else 2.0
            cw = (W - 1.2 - gap * (per - 1)) / per
            ch = min(4.0, (6.8 - top - gap * (rows - 1)) / rows)
            for j, c in enumerate(cards):
                card(s, 0.6 + (j % per) * (cw + gap), top + (j // per) * (ch + gap), cw, ch,
                     c["title"], c.get("body", ""), c.get("highlight", False), c.get("tag"), c.get("icon"))
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
    prs.save(out)


# ───────────────────────── PDF ─────────────────────────
def build_pdf(spec, out):
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    PW, PH = W * 72, H * 72
    reg, bold = "Helvetica", "Helvetica-Bold"
    for r, b in [("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
                 ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf")]:
        if os.path.exists(r) and os.path.exists(b):
            pdfmetrics.registerFont(TTFont("TH", r)); pdfmetrics.registerFont(TTFont("TH-B", b)); reg, bold = "TH", "TH-B"; break
    cv = canvas.Canvas(out, pagesize=(PW, PH))
    footer_text = spec.get("footer", FOOTER)
    hx = lambda n: HexColor("#" + col(n))
    Y = lambda y, h=0: PH - (y + h) * 72  # top-based inches -> pdf coords

    def rect(x, y, w, h, fill, r=0, stroke=None):
        cv.setFillColor(hx(fill))
        if stroke: cv.setStrokeColor(hx(stroke)); cv.setLineWidth(0.75)
        (cv.roundRect(x*72, Y(y, h), w*72, h*72, r*72, stroke=1 if stroke else 0, fill=1) if r
         else cv.rect(x*72, Y(y, h), w*72, h*72, stroke=1 if stroke else 0, fill=1))

    def circ(cx, cy, d, fill):
        cv.setFillColor(hx(fill)); cv.circle(cx*72, PH - cy*72, d/2*72, stroke=0, fill=1)

    def wrap(t, font, size, maxw):
        out_l = []
        for para in str(t).split("\n"):
            line = ""
            for wd in para.split():
                trial = (line + " " + wd).strip()
                if pdfmetrics.stringWidth(trial, font, size) <= maxw * 72: line = trial
                else: out_l.append(line); line = wd
            out_l.append(line)
        return out_l

    def txt(x, y, w, t, size, color, b=False, align="l", lead=1.25, maxlines=None):
        font = bold if b else reg
        cv.setFillColor(hx(color)); cv.setFont(font, size)
        lines = wrap(t, font, size, w)[:maxlines]
        for k, ln in enumerate(lines):
            by = PH - y*72 - size - k * size * lead
            if align == "c": cv.drawCentredString((x + w/2)*72, by, ln)
            elif align == "r": cv.drawRightString((x + w)*72, by, ln)
            else: cv.drawString(x*72, by, ln)
        return len(lines)

    def pill(x, y, label, align_right=False):
        w = max(0.9, 0.095 * len(label) + 0.45)
        if align_right: x = x - w
        rect(x, y, w, 0.33, badge_color(label), r=0.165)
        txt(x, y + 0.085, w, label.upper(), 10, "white", True, "c")

    for i, sl in enumerate(spec["slides"], 1):
        k = sl.get("type", "cards")
        if k in ("cover", "section"):
            rect(0, 0, W, H, "bg_dark")
            circ(9.4 + 2.8, -2.2 + 2.8, 5.6, "circle_a")
            if k == "cover":
                circ(10.2 + 2.3, 3.9 + 2.3, 4.6, "circle_b")
                if sl.get("arabic"): txt(0.6, 1.35, 8, sl["arabic"], 26, "gold", True)
                txt(0.6, 1.95, 9, sl["title"], 50, "white", True)
                txt(0.6, 3.0, 9, sl.get("subtitle", ""), 28, "lilac", True)
                txt(0.6, 3.85, 8.5, sl.get("tagline", ""), 16, "lilac_soft")
                txt(0.6, 5.6, 9, sl.get("footnote", ""), 12, "lilac_dim")
            else:
                txt(0.6, 2.6, 9, sl.get("eyebrow", "").upper(), 13, "gold", True)
                txt(0.6, 3.0, 10, sl["title"], 42, "white", True)
                if sl.get("lede"): txt(0.6, 4.3, 9, sl["lede"], 18, "lilac")
        else:
            rect(0, 0, W, H, "bg_light")
            txt(0.6, 0.45, 8, sl.get("eyebrow", "").upper(), 12, "green", True)
            txt(0.6, 0.8, 10.5, sl["title"], 34, "ink", True)
            if sl.get("lede"): txt(0.6, 1.7, 11.5, sl["lede"], 15, "muted")
            if sl.get("badge"): pill(W - 0.6, 1.0, sl["badge"], True)
            txt(0.6, 7.07, 6, footer_text, 10, "muted")
            txt(W - 1.6, 7.07, 1.0, str(i), 10, "muted", align="r")
            if k == "bullets":
                y = 2.4
                for b in sl["bullets"]:
                    circ(0.73, y + 0.18, 0.16, "green"); txt(1.0, y, 11, b, 17, "body"); y += 0.62
            else:
                cards = sl["cards"]; n = len(cards); per = sl.get("per_row", min(n, 4)); rows = -(-n // per)
                gap = 0.3; top = 2.3 if sl.get("lede") else 2.0
                cw = (W - 1.2 - gap * (per - 1)) / per
                ch = min(4.0, (6.8 - top - gap * (rows - 1)) / rows)
                for j, c in enumerate(cards):
                    x = 0.6 + (j % per) * (cw + gap); y = top + (j // per) * (ch + gap)
                    rect(x, y, cw, ch, "green_tint" if c.get("highlight") else "card", r=0.15, stroke="line")
                    circ(x + 0.22 + 0.31, y + 0.22 + 0.31, 0.62, "green_soft")
                    if c.get("icon"): txt(x + 0.22, y + 0.22 + 0.17, 0.62, c["icon"], 16, "green", True, "c")
                    if c.get("tag"):
                        pw = max(0.6, 0.085 * len(c["tag"]) + 0.35)
                        rect(x + cw - pw - 0.18, y + 0.24, pw, 0.27, badge_color(c["tag"]), r=0.135)
                        txt(x + cw - pw - 0.18, y + 0.285, pw, c["tag"].upper(), 9, "white", True, "c")
                    txt(x + 0.22, y + 0.98, cw - 0.44, c["title"], 14, "ink", True)
                    txt(x + 0.22, y + 1.42, cw - 0.44, c.get("body", ""), 11.5, "body")
        cv.showPage()
    cv.save()


# ───────────────────────── RESTYLE ─────────────────────────
def restyle(src, out, dark_first=True):
    """Apply Aqarati palette to an existing PPTX: backgrounds, fonts, text colors, footers."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt, Inches
    prs = Presentation(src)
    rgb = lambda n: RGBColor.from_string(col(n))
    for idx, s in enumerate(prs.slides):
        dark = dark_first and idx == 0
        f = s.background.fill; f.solid(); f.fore_color.rgb = rgb("bg_dark" if dark else "bg_light")
        for sh in s.shapes:
            if not sh.has_text_frame: continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = F["main"]
                    size = r.font.size.pt if r.font.size else 14
                    if dark: r.font.color.rgb = rgb("white" if size >= 30 else "lilac")
                    elif size >= 30: r.font.color.rgb = rgb("ink"); r.font.bold = True
                    elif size >= 20: r.font.color.rgb = rgb("ink")
                    else: r.font.color.rgb = rgb("body")
        if not dark:
            tb = s.shapes.add_textbox(Inches(0.6), Inches(7.07), Inches(6), Inches(0.3))
            tb.text_frame.text = FOOTER
            r = tb.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(10); r.font.name = F["main"]; r.font.color.rgb = rgb("muted")
    prs.save(out)


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) < 3 or a[0] not in ("build", "restyle"):
        sys.exit(__doc__)
    if a[0] == "restyle":
        restyle(a[1], a[2]); print("wrote", a[2])
    else:
        spec = json.load(open(a[1])); out = a[2]
        if out.endswith(".pdf"): build_pdf(spec, out)
        else:
            build_pptx(spec, out)
            if "--pdf" in a: build_pdf(spec, out[:-5] + ".pdf")
        print("wrote", out)
